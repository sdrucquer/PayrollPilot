import copy
import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from payrollpilot_core.audit import run_audit
from payrollpilot_core.corrections import apply_corrections, propose_correction
from payrollpilot_core.packet import build_packet
from payrollpilot_core.workbook import PENDING, SHEETS, WorkbookVerificationError, build_workbook, verify_workbook

FIX = Path(__file__).resolve().parents[2] / "docs" / "contracts" / "fixtures"
REQUIRED = ["Review", "Labor", "Time", "Manual payments", "Payments", "Expenses", "Overhead", "Quotes & marketing",
            "Tips", "Questions", "Job queue", "Carry-forward"]


@pytest.fixture()
def snap():
    return json.loads((FIX / "cycle_synthetic_v1.json").read_text())


@pytest.fixture()
def packet(snap):
    return build_packet(snapshot=snap, exceptions=run_audit(snap))


def test_packet_contents(packet):
    assert packet["calculation_version"] == "calc_v1" and packet["destination"] == "internal_draft"
    assert packet["cycle"]["id"] == "cyc_synth_v1" and packet["workbook_ref"] is None
    t = packet["totals"]
    assert t["workers_review_ready"] == 1 and t["stage1_gross_ready_cents"] == 174900
    assert t["stage1_gross_all_cents"] is None  # unknown, never summed as zero
    assert t["eligible_receipts_cents"] == 450000 and t["tips_cents"] == 5000
    assert t["entries_excluded_out_of_cycle"] == 1 and t["cycle_profit_cents"] is None
    assert packet["readiness"]["stage1"]["state"] == "partial"
    assert {i["id"] for i in packet["readiness"]["items"]} >= {"ci_quote_job_01", "ci_manager_pool"}
    assert all(set(i) == {"worker_id", "display_name", "identity_status", "payment_route"} for i in packet["identities"])


def test_draft_version_stable_and_changes_with_inputs(snap, packet):
    assert build_packet(snapshot=copy.deepcopy(snap), exceptions=run_audit(snap))["draft_version"] == packet["draft_version"]
    corr = propose_correction("job_allocation", {"time_entry_id": "te_0014", "job_id": "job_01"},
                              {"time_entry_id": "te_0014", "job_id": "job_02"}, snap)
    view = apply_corrections(snap, [corr])
    p2 = build_packet(snapshot=view, exceptions=run_audit(view), corrections=[corr])
    assert p2["draft_version"] != packet["draft_version"]
    assert p2["totals"]["stage1_gross_ready_cents"] == packet["totals"]["stage1_gross_ready_cents"]  # T02: wages same


def test_t28_workbook_all_ledgers_verified(packet, tmp_path):
    art = build_workbook(packet, tmp_path / "review.xlsx")
    assert art["original_updated"] is False and art["verified"] is True
    assert art["note"] == "Provisional review workbook; designated workbook not supplied/updated."
    assert art["draft_version"] == packet["draft_version"]
    assert art["verified_totals"] == packet["totals"]
    wb = load_workbook(art["path"])
    assert wb.sheetnames == REQUIRED + ["Sources"] == SHEETS


def test_t28_pending_tabs_present_and_labeled(packet, tmp_path):
    wb = load_workbook(build_workbook(packet, tmp_path / "r.xlsx")["path"])
    for name in ("Overhead", "Quotes & marketing", "Tips", "Manual payments", "Carry-forward"):
        texts = [str(c.value) for row in wb[name].iter_rows() for c in row if c.value is not None]
        assert any(t.startswith(PENDING) for t in texts), name
        assert wb[name].max_row >= 2  # header + status row
    overhead = [str(c.value) for row in wb["Overhead"].iter_rows() for c in row if c.value]
    assert any("O05" in t for t in overhead)


def test_no_real_names_or_contact_data(packet, tmp_path):
    wb = load_workbook(build_workbook(packet, tmp_path / "r.xlsx")["path"])
    phone = re.compile(r"\+\d{10,15}|\+1[\s.-]?\d{3}[\s.-]?\d{3}[\s.-]?\d{4}|\(\d{3}\)\s?\d{3}-\d{4}|\b\d{3}[.-]\d{3}[.-]\d{4}\b")
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str):
                    assert "@" not in v and not phone.search(v), (ws.title, v)
    names = {r[0] for r in wb["Labor"].iter_rows(min_row=2, values_only=True)}
    assert names <= {"Worker A", "Worker B", "Worker C", "Worker D"}


def test_tampered_workbook_fails_verification(packet, tmp_path):
    art = build_workbook(packet, tmp_path / "r.xlsx")
    wb = load_workbook(art["path"])
    for row in wb["Review"].iter_rows(min_row=2):
        if row[0].value == "total:stage1_gross_ready_cents":
            row[1].value = 999
    wb.save(art["path"])
    with pytest.raises(WorkbookVerificationError):
        verify_workbook(art["path"], packet)


def test_legacy_positional_packet_reports_unknown_not_zero(snap, tmp_path):
    excs = run_audit(snap)
    p = build_packet(snap["cycle"], snap["workers"], snap["time_entries"], [], excs, [])
    assert p["partial_inputs"] is True and p["totals"]["eligible_receipts_cents"] is None
    assert p["totals"]["tips_cents"] is None
    assert build_workbook(p, tmp_path / "legacy.xlsx")["verified"] is True
