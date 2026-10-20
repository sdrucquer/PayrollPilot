"""Reconciliation workbook (FR-28, T28). Writes 12 ledgers + Sources to a caller-supplied path, reopens the
file, and verifies the Review totals equal the packet's totals for the same draft version.

Pending ledgers keep a header row and a "status: pending — <blocker>" row. The designated original workbook
is never touched: original_updated is always False in this pilot.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from . import CALC_VERSION
from .ids import content_hash

# Excel forbids "/" in sheet names, so the PRD's "Quotes/marketing" ledger is named "Quotes & marketing".
SHEETS = ["Review", "Labor", "Time", "Manual payments", "Payments", "Expenses", "Overhead", "Quotes & marketing",
          "Tips", "Questions", "Job queue", "Carry-forward", "Sources"]
NOTE = "Provisional review workbook; designated workbook not supplied/updated."
PENDING = "status: pending — "


class WorkbookVerificationError(RuntimeError):
    pass


def _cell(v):
    if v is None:
        return "unknown"
    if isinstance(v, (list, dict)):
        return str(v)
    return v


def _write(ws, header: list[str], rows: list[list], pending: str | None = None) -> None:
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(v) for v in r])
    if pending:
        ws.append([PENDING + pending])
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 70)


def _blockers_text(readiness: dict) -> str:
    return "; ".join(f"{b['code']} ({b['label']})" for b in readiness.get("blockers", [])) or "none"


def build_workbook(packet: dict, path: str | Path) -> dict:
    """Write the workbook, reopen and verify it, return a WorkbookArtifact dict."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    ws = {name: wb.create_sheet(name) for name in SHEETS}
    totals = packet["totals"]
    rd = packet["readiness"]
    items = rd["items"]

    # Review
    review_rows = [["draft_version", packet["draft_version"]], ["calculation_version", packet["calculation_version"]],
                   ["cycle", packet["cycle"]["id"]], ["company", packet["cycle"]["company_label"]],
                   ["period", f"{packet['cycle']['period_start']} .. {packet['cycle']['period_end']}"],
                   ["destination", packet["destination"]], ["destination_note", packet["destination_note"]],
                   ["provenance", packet["provenance"]], ["original_updated", "false"], ["note", NOTE],
                   ["stage1", f"{rd['stage1']['state']}: {rd['stage1']['scope']}"]]
    review_rows += [[f"item:{i['id']}", f"{i['readiness']['state']}: {i['readiness']['scope']}"] for i in items]
    review_rows += [[f"total:{k}", v] for k, v in totals.items()]
    _write(ws["Review"], ["field", "value (money in cents; unknown = not computable)"], review_rows)
    for row in ws["Review"].iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(horizontal="left")

    # Labor
    labor = []
    for e in packet["earnings"]:
        for week, minutes in e["weekly_minutes"].items():
            labor.append([e["display_name"], e["worker_id"], week, minutes,
                          ", ".join(f"{r['rate_cents_per_hour'] if r and r['rate_cents_per_hour'] is not None else 'unconfirmed'}"
                                    f"/{r['category'] if r else 'unknown'}" for r in e["rates"]),
                          e["gross_new_cents"], e["readiness"]["state"], _blockers_text(e["readiness"])])
    _write(ws["Labor"], ["worker", "worker_id", "workweek_start", "minutes", "rate cents/h / category",
                         "gross_new_cents (cycle)", "readiness", "blockers"], labor,
           None if labor else "no in-cycle time")

    # Time
    period = (packet["cycle"]["period_start"], packet["cycle"]["period_end"])
    trow = [[t["id"], t["worker_id"], t.get("job_id") or "none", t["date"], t.get("start"), t.get("end"),
             t.get("duration_minutes"), "yes" if period[0] <= t["date"] <= period[1] else "excluded (out of cycle)",
             t.get("paid_status"), t.get("provenance")] for t in packet["time_entries"]]
    _write(ws["Time"], ["entry", "worker_id", "job_id", "date", "start", "end", "minutes", "in_cycle", "paid_status",
                        "provenance"], trow, None if trow else "time entries not supplied")

    # Manual payments
    manual = [[e["display_name"], e["payment_route"], e["gross_new_cents"]] for e in packet["earnings"]
              if e["payment_route"] in ("manual", "unknown")]
    _write(ws["Manual payments"], ["worker", "payment_route", "gross_new_cents"], manual,
           ("payment route unconfirmed for " + ", ".join(m[0] for m in manual)) if manual
           else "no manual payment records supplied")

    # Payments
    prow = [[r["id"], r["job_id"], r["kind"], r["amount_cents"], r["payment_date"], r["status"],
             (r.get("cycle_allocation") or {}).get("state"), (r.get("cycle_allocation") or {}).get("allocated_cycle_id")]
            for r in packet["receipts"]]
    _write(ws["Payments"], ["receipt", "job_id", "kind", "amount_cents", "payment_date", "status", "allocation_state",
                            "allocated_cycle"], prow, None if prow else "receipts not supplied")

    # Expenses
    erow = [[c["id"], c["description"], c.get("amount_cents"), c.get("cost_date"),
             ", ".join(f"{r['kind']}:{r.get('amount_cents')}" for r in c.get("representations", [])),
             ", ".join(f"{a['job_id']}:{a['amount_cents']}" for a in c.get("job_allocations", [])) or "none",
             c.get("categorized"), c.get("allocated"), c.get("included_in_calc"), c.get("conflict") or "none"]
            for c in packet["costs"]]
    _write(ws["Expenses"], ["cost (one economic identity)", "description", "amount_cents", "date", "representations",
                            "job_allocations", "categorized", "allocated", "included", "conflict"], erow,
           None if erow else "costs not supplied")

    # Overhead: always pending in this pilot
    mgr = next((i for i in items if i["kind"] == "manager_pool"), None)
    oh = next((b["label"] for b in (mgr or {}).get("readiness", {}).get("blockers", []) if b["code"] == "policy_missing:O05"),
              "overhead schedule not supplied")
    _write(ws["Overhead"], ["item", "amount_cents", "basis", "status"], [], oh)

    # Quotes & marketing
    qrow = [[i["id"], i["kind"], i.get("job_id"), (i.get("computed") or {}).get("pool_cents"),
             str((i.get("computed") or {}).get("shares")) if i.get("computed") else "unknown",
             i["readiness"]["state"], _blockers_text(i["readiness"])]
            for i in items if i["kind"] in ("quote_commission", "marketing")]
    pend = [r for r in qrow if r[5] != "ready"]
    _write(ws["Quotes & marketing"], ["item", "kind", "job_id", "pool_cents", "shares_cents", "readiness", "blockers"],
           qrow, f"{len(pend)} item(s) not ready" if pend or not qrow else None)

    # Tips
    tip_items = [i for i in items if i["kind"] == "tip"]
    tirow = [[i["id"], i.get("job_id"), i["inputs"].get("tip_cents"), i["readiness"]["state"], _blockers_text(i["readiness"])]
             for i in tip_items]
    _write(ws["Tips"], ["item", "job_id", "tip_cents (excluded from company receipts)", "readiness", "blockers"], tirow,
           "tip allocation unresolved" if any(r[3] != "ready" for r in tirow) else (None if tirow else "no tips"))

    # Questions
    qs = [[e["id"], e["category"], e["title"], e["intended_respondent_role"], e.get("state"), e.get("next_action"),
           e["stage_impact"]["stage1_wages"], e["stage_impact"]["stage2_costs"], e.get("dollars_affected_cents"),
           e.get("provenance")] for e in packet["exceptions"]]
    _write(ws["Questions"], ["exception", "category", "title", "intended_respondent", "state", "next_action",
                             "stage1_impact", "stage2_impact", "dollars_affected_cents", "provenance"], qs)

    # Job queue
    jrow = [[j["job_id"], j["display_name"], j["status"], j["adjusted_invoice_cents"], j["final_payment_verified"],
             j["labor_cents_this_cycle"], j["included_costs_cents"], j["open_exceptions"]] for j in packet["jobs"]]
    _write(ws["Job queue"], ["job", "name", "status", "adjusted_invoice_cents", "final_payment_verified",
                             "labor_cents_this_cycle", "included_costs_cents", "open_exceptions"], jrow,
           None if jrow else "jobs not supplied")

    # Carry-forward
    crow = [[o["id"], o["kind"], o.get("worker_id"), o.get("recognition_cycle_id"), o.get("gross_cents"),
             o.get("balance_cents"), o.get("status")] for o in packet["obligations"]]
    _write(ws["Carry-forward"], ["obligation", "kind", "worker_id", "recognition_cycle", "gross_cents", "balance_cents",
                                 "status"], crow,
           "manager pool distributions blocked: " + _blockers_text(mgr["readiness"]) if mgr else None)

    # Sources
    ev = packet["evidence"]
    srows = [["snapshot_id", ev.get("snapshot_id")], ["captured_at", ev.get("captured_at")],
             ["draft_version", packet["draft_version"]], ["calculation_version", CALC_VERSION],
             ["contract_version", packet["contract_version"]], ["policy_version", packet["cycle"].get("policy_version")]]
    srows += [[f"coverage:{c['system']}:{c['dataset']}",
               f"{c['record_count']} records; complete={c['complete']}; method={c['method']}; reason={c.get('incomplete_reason')}"]
              for c in ev.get("coverage", [])]
    _write(ws["Sources"], ["source", "detail"], srows)

    wb.save(path)
    verified = verify_workbook(path, packet)
    return {"id": "wb_" + content_hash([packet["draft_version"], str(path.name)])[7:17], "path": str(path),
            "draft_version": packet["draft_version"], "sheets": SHEETS, "verified_totals": verified,
            "verified": True, "original_updated": False, "note": NOTE, "calculation_version": CALC_VERSION}


def verify_workbook(path: str | Path, packet: dict) -> dict:
    """Reopen the file; assert all sheets exist and Review totals + draft version equal the packet's."""
    wb = load_workbook(path, read_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise WorkbookVerificationError(f"missing sheets: {missing}")
    review = {row[0]: row[1] for row in wb["Review"].iter_rows(min_row=2, values_only=True) if row and row[0]}
    if review.get("draft_version") != packet["draft_version"]:
        raise WorkbookVerificationError("draft_version mismatch")
    out = {}
    for k, v in packet["totals"].items():
        got = review.get(f"total:{k}")
        want = "unknown" if v is None else v
        if got != want:
            raise WorkbookVerificationError(f"total {k}: workbook {got!r} != packet {want!r}")
        out[k] = v
    for name in SHEETS:
        if wb[name].max_row < 1:
            raise WorkbookVerificationError(f"sheet {name} empty")
    wb.close()
    return out
